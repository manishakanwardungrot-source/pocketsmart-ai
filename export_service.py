from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


EXPORT_DIR = Path("exports")
EXPORT_DIR.mkdir(parents=True, exist_ok=True)


def _safe_filename(filename: str) -> str:
    filename = Path(filename).stem

    safe_name = "".join(
        char if char.isalnum() or char in (" ", "_", "-") else "_"
        for char in filename
    )

    return safe_name.strip() or "document"


def export_to_txt(content: str, filename: str = "document.txt") -> str:
    safe_name = _safe_filename(filename)
    output_path = EXPORT_DIR / f"{safe_name}.txt"

    output_path.write_text(content, encoding="utf-8")

    return str(output_path)


def export_to_docx(content: str, filename: str = "document.docx") -> str:
    safe_name = _safe_filename(filename)
    output_path = EXPORT_DIR / f"{safe_name}.docx"

    document = Document()
    document.add_heading(safe_name, level=1)

    for paragraph in content.split("\n"):
        document.add_paragraph(paragraph)

    document.save(output_path)

    return str(output_path)


def export_to_pdf(content: str, filename: str = "document.pdf") -> str:
    safe_name = _safe_filename(filename)
    output_path = EXPORT_DIR / f"{safe_name}.pdf"

    pdf = canvas.Canvas(str(output_path), pagesize=A4)

    width, height = A4
    x = 50
    y = height - 50
    line_height = 16

    pdf.setFont("Helvetica", 11)

    for paragraph in content.split("\n"):
        words = paragraph.split()
        current_line = ""

        for word in words:
            test_line = (
                f"{current_line} {word}"
                if current_line
                else word
            )

            if pdf.stringWidth(test_line, "Helvetica", 11) > width - 100:
                if current_line:
                    pdf.drawString(x, y, current_line)

                y -= line_height
                current_line = word

                if y < 50:
                    pdf.showPage()
                    pdf.setFont("Helvetica", 11)
                    y = height - 50
            else:
                current_line = test_line

        if current_line:
            pdf.drawString(x, y, current_line)
            y -= line_height

        if y < 50:
            pdf.showPage()
            pdf.setFont("Helvetica", 11)
            y = height - 50

    pdf.save()

    return str(output_path)


def export_document(
    content: str,
    filename: str,
    file_format: str
) -> str:

    file_format = file_format.lower().strip()

    if file_format == "txt":
        return export_to_txt(content, filename)

    if file_format == "docx":
        return export_to_docx(content, filename)

    if file_format == "pdf":
        return export_to_pdf(content, filename)

    raise ValueError(
        f"Unsupported export format: {file_format}"
    )
